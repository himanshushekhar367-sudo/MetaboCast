#!/usr/bin/env python3
"""Build the Supplementary Tables workbook (Additional file 2) from reproduce/results.

Table numbers and titles come from reproduce/manifest.py (one sheet per table, named S1, S2, ...).

Usage:
    python reproduce/tables/build_supplementary_tables.py [--results reproduce/results] [--out MetaboCast_Supplementary_Tables.xlsx]
Also writes <out>_notes.json with the summary values quoted in the manuscript.
"""
import argparse
import json
import os
import re
import sys

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.stats import spearmanr, wilcoxon

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from manifest import TABLES, TAB, table_sheet, table_ref as TR, figure_ref as FR          # noqa: E402

DESIUM = ["BC_515_Section_1", "BC_515_Section_2", "BC_525", "BC_823", "LC_091", "LC_170", "LC_276"]
CANCER = ["breast (TNBC)", "breast (TNBC)", "breast", "breast", "lung", "lung", "lung"]
PATIENT = {"BC_515_Section_1": "BC_515", "BC_515_Section_2": "BC_515"}
SMA_GROUPS = {"dhb": ["SMA_DHB_V11L12-038_A1", "SMA_DHB_V11L12-038_B1"],
              "fmp_mouse": ["SMA_FMP_V11L12-109_A1", "SMA_FMP_V11L12-109_B1", "SMA_FMP_V11L12-109_C1",
                            "SMA_FMP_V11T16-085_A1", "SMA_FMP_V11T16-085_B1", "SMA_FMP_V11T16-085_C1"],
              "fmp_human": ["SMA_FMP_V11T17-102_A1", "SMA_FMP_V11T17-102_B1", "SMA_FMP_V11T17-102_C1", "SMA_FMP_V11T17-102_D1"]}
NIGRA = {"SMA_FMP_V11T16-085_A1", "SMA_FMP_V11T16-085_B1", "SMA_FMP_V11T16-085_C1"}
EXPECTED = {"Glutathione(GSH)": 1, "Ascorbate": 1, "Arachidonate(FA20:4)": 1, "DHA(FA22:6)": 1,
            "Palmitate(FA16:0)": -1, "Stearate(FA18:0)": -1, "Hexose": -1}
PATIENTS = ["HCC-1", "HCC-2", "HCC-3", "HCC-4"]
KEGG = "KEGG pathway features + supervised regressor"
COORD = "coordinates only (tuned kNN)"
TRANSFER_MIN, WITHIN_MIN = 0.20, 0.30
HFILL = PatternFill("solid", fgColor="DDE6F0")


class Book:
    def __init__(self):
        self.wb = openpyxl.Workbook(); self.wb.remove(self.wb.active); self.notes = {}; self.b_titles = {}

    def sheet(self, key, header, rows, notes=(), widths=None):
        name = table_sheet(key); title = f"{TAB[key]['label']} | {TAB[key]['title']}"
        ws = self.wb.create_sheet(name)
        ws.append([title]); ws["A1"].font = Font(bold=True, size=11)
        ws.append([]); ws.append(header)
        for c in ws[3]:
            c.font = Font(bold=True); c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([clean(x) for x in r])
        if notes:
            ws.append([])
            for n in notes:
                ws.append([n]); ws.cell(ws.max_row, 1).font = Font(italic=True, size=9)
        for i in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths and i <= len(widths)
                                                                 else max(12, min(40, len(str(header[i - 1])) + 4)))
        ws.freeze_panes = "A4"
        return ws

    def sheet_b(self, key, title, header, rows, notes=(), widths=None):
        """Sheet b of a numbered table (e.g. Table S5b), written like sheet()."""
        name = table_sheet(key) + "b"
        ws = self.wb.create_sheet(name)
        ws.append([f"Table {name} | {title}"]); ws["A1"].font = Font(bold=True, size=11)
        ws.append([]); ws.append(header)
        for c in ws[3]:
            c.font = Font(bold=True); c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([clean(x) for x in r])
        if notes:
            ws.append([])
            for n in notes:
                ws.append([n]); ws.cell(ws.max_row, 1).font = Font(italic=True, size=9)
        for i in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths and i <= len(widths) else 14)
        ws.freeze_panes = "A4"
        self.b_titles[name] = title
        return ws


def clean(x):
    if isinstance(x, (float, np.floating)):
        return None if np.isnan(x) else round(float(x), 4)
    if isinstance(x, np.integer):
        return int(x)
    if isinstance(x, np.bool_):
        return bool(x)
    return x


def msd(v):
    v = np.asarray([x for x in v if x == x], float)
    if not len(v):
        return "n/a"
    return f"{v.mean():.3f}" + (f" ± {v.std(ddof=1):.3f}" if len(v) > 1 else "")


def boot_ci(a, b, n=5000, seed=0):
    d = np.asarray(a, float) - np.asarray(b, float); r = np.random.RandomState(seed)
    m = d[r.randint(0, len(d), (n, len(d)))].mean(1)
    return float(d.mean()), float(np.percentile(m, 2.5)), float(np.percentile(m, 97.5))


def wp(d):
    d = np.asarray(d, float)
    return float(wilcoxon(d).pvalue) if np.any(d != 0) else 1.0


def boot_median(d, n=5000, seed=0):
    """Median of paired differences with a percentile bootstrap interval (resampling the paired units)."""
    d = np.asarray(d, float); r = np.random.RandomState(seed)
    m = np.median(d[r.randint(0, len(d), (n, len(d)))], axis=1)
    return float(np.median(d)), float(np.percentile(m, 2.5)), float(np.percentile(m, 97.5))


def holm(ps):
    """Holm step-down adjustment of a family of P values (family-wise error rate)."""
    ps = np.asarray(ps, float); adj = np.empty(len(ps)); run = 0.0
    for k, i in enumerate(np.argsort(ps)):
        run = max(run, (len(ps) - k) * ps[i]); adj[i] = min(1.0, run)
    return adj


def pw(a):
    """A penalty as a power of ten for display: 1e7 -> 10⁷, 3e8 -> 3 × 10⁸."""
    sup = str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")
    m, e = f"{float(a):.0e}".split("e")
    e = str(int(e)).translate(sup)
    return f"10{e}" if m == "1" else f"{m} × 10{e}"


def tier(w, t):
    if t == t and t >= TRANSFER_MIN:
        return "transfer-supported"
    if w == w and w >= WITHIN_MIN:
        return "within-section only"
    return "low" if (w == w or t == t) else "uncalibrated"


def effect_rows(names, x_of, pairs, prefix=""):
    rows, out = [], {}
    for a, b in pairs:
        x, y = x_of(a), x_of(b)
        if np.isnan(x).any() or np.isnan(y).any():
            continue
        m_, lo_, hi_ = boot_ci(x, y); dd = x - y; p = wp(dd)
        rows.append([prefix + f"{a} − {b}", f"{m_:+.4f}", f"[{lo_:+.4f}, {hi_:+.4f}]", f"{int((dd > 0).sum())}/{len(dd)}",
                     round(p, 4), "yes" if (lo_ > 0 or hi_ < 0) else "no"])
        out[f"{a} - {b}"] = dict(delta=m_, ci=(lo_, hi_), wins=int((dd > 0).sum()), n=len(dd), p=p)
    return rows, out


def build(R, out):
    import warnings                                    # channels never detected give empty slices (NaN), as intended
    warnings.filterwarnings("ignore", message="Mean of empty slice")
    warnings.filterwarnings("ignore", message="Degrees of freedom <= 0 for slice")
    J = lambda *p: json.load(open(os.path.join(R, *p)))
    med = lambda *p: {s: v["median"] for s, v in J(*p)["per_section"].items()}
    X = lambda f: pd.read_csv(os.path.join(R, "crosscancer", f))
    A = lambda f: pd.read_csv(os.path.join(R, "ablations", f))
    B = Book(); N = B.notes

    # ------------------------------------------------------------- benchmark (within-section, DESIUM)
    wj = J("within_desium.json"); ctrl = J("within_controls.json"); secs = [s for s in DESIUM if s in wj]

    def sc(m, t=False):
        if m == COORD:
            return np.array([np.nan if t else ctrl[s]["pooled"]["coordinates only"]["full_panel"] for s in secs])
        return np.array([(wj[s]["scores_target_scale"] if t else wj[s]["scores"]).get(m, np.nan) for s in secs])
    meths = ["MetaboCast", "MetaboCast + KEGG", "MetaboCast (ridge only)", "no diffusion", KEGG, "joint-NMF",
             COORD, "single-gene OLS"]
    pairs = [("MetaboCast", "joint-NMF"), ("MetaboCast", COORD), ("MetaboCast", "no diffusion"), ("MetaboCast", KEGG),
             ("MetaboCast", "MetaboCast (ridge only)"), ("MetaboCast + KEGG", "MetaboCast"), ("MetaboCast", "single-gene OLS")]
    rows, N["benchmark"] = effect_rows(secs, sc, pairs)
    r2, N["benchmark_target"] = effect_rows(secs, lambda m: sc(m, True), pairs, prefix="target scale: ")
    rows += r2 + [[], ["Per-method mean ± s.d. across 7 sections", "intensity scale", "target scale"]]
    rows += [[m, msd(sc(m)), msd(sc(m, True))] for m in meths]
    N["within"] = {m: [float(np.nanmean(sc(m))), float(np.nanstd(sc(m), ddof=1))] for m in meths}
    B.sheet("benchmark", ["Comparison", "Δ median Spearman", "95% CI (paired bootstrap, 5,000)", "Sections improved",
                          "Wilcoxon P (exact)", "CI excludes 0"], rows,
            ["Five k-means spatial blocks per section; highly variable genes, PCA, penalty search and both regressors use the training "
             "blocks only; out-of-fold predictions pooled per section; median per-ion Spearman over the 1,000 most variable ions detected "
             "in >20% of spots (reproduce/within_benchmark.py). Exact two-sided Wilcoxon over 7 sections (smallest attainable P = 0.016).",
             "Regression baselines use the MetaboCast estimator with different features: no diffusion, PCA features without the diffused copies; "
             "KEGG pathway features + supervised regressor, mean z-scored member-gene expression of 85 KEGG metabolic gene sets (as "
             "distributed with scMetabolism) and one diffused copy, fitted to the measured MSI (scMetabolism itself does not predict "
             "metabolites); ridge only, without the tree half. Joint NMF and single-gene OLS are fitted on the training blocks of the panel.",
             "Coordinates only: k-nearest-neighbour regression on spot coordinates from the measured metabolites of nearby training spots; "
             "k (1-320) and the weighting (uniform or inverse distance) are chosen per fold by leave-one-block-out CV inside the training "
             f"blocks (reproduce/within_controls.py; {TR('within_controls')}). It has no target scale.",
             "'target scale' rows pool each fold model's predictions on its own training-target scale (rank-Gaussian scores), a "
             "sensitivity analysis."], [56, 16, 26, 14, 16, 14])

    # ------------------------------------------------------------- within-section controls
    rows = []
    for s in secs:
        c = ctrl[s]; pf = pd.DataFrame(c["per_fold"])
        rows.append([s, c["n_spots"], c["n_detected"], "; ".join(f"k={x['k']}{' (1/d)' if x['weighted'] else ''}" for x in c["knn_choice"]),
                     c["pooled"]["MetaboCast"]["full_panel"], c["pooled"]["coordinates only"]["full_panel"],
                     c["pooled"]["MetaboCast"]["all_detected"], c["pooled"]["coordinates only"]["all_detected"],
                     pf.metabocast_full_panel.mean(), pf.metabocast_train_panel.mean(), pf.coords_full_panel.mean(),
                     pf.coords_train_panel.mean(), c["n_panel_every_fold"], c["pooled"]["MetaboCast"]["every_fold_panel"],
                     c["pooled"]["coordinates only"]["every_fold_panel"]])
    sm = ctrl["_summary"]
    rows.append(["mean over sections", "", "", "", sm["pooled_mean_over_sections"]["MetaboCast"]["full_panel"],
                 sm["pooled_mean_over_sections"]["coordinates only"]["full_panel"],
                 sm["pooled_mean_over_sections"]["MetaboCast"]["all_detected"],
                 sm["pooled_mean_over_sections"]["coordinates only"]["all_detected"],
                 sm["per_block_mean_over_folds"]["metabocast_full_panel"], sm["per_block_mean_over_folds"]["metabocast_train_panel"],
                 sm["per_block_mean_over_folds"]["coords_full_panel"], sm["per_block_mean_over_folds"]["coords_train_panel"], "",
                 sm["pooled_mean_over_sections"]["MetaboCast"]["every_fold_panel"],
                 sm["pooled_mean_over_sections"]["coordinates only"]["every_fold_panel"]])
    N["within_controls"] = sm
    B.sheet("within_controls",
            ["Section", "Spots", "Ions detected (>20%)", "kNN choice per fold", "MetaboCast, panel", "Coordinates, panel",
             "MetaboCast, all detected", "Coordinates, all detected", "MetaboCast per block, whole-section panel",
             "MetaboCast per block, training-block panel", "Coordinates per block, whole-section panel",
             "Coordinates per block, training-block panel", "Ions in every fold's panel", "MetaboCast, every-fold panel",
             "Coordinates, every-fold panel"], rows,
            [f"Same sections, blocks and MetaboCast fits as {TR('benchmark')} (reproduce/within_controls.py). 'Panel' columns pool the out-of-fold "
             "predictions as in the benchmark; 'per block' columns score each held-out block separately and average over the five blocks.",
             "Ion-panel selection: the benchmark panel (1,000 most variable ions detected in >20% of spots) is chosen on the whole section. "
             "The training-block panel repeats the choice on the training blocks of each fold only; ions chosen in every fold are also "
             "scored with the pooled predictions.",
             "kNN choice: k and weighting (1/d = inverse distance) chosen by leave-one-block-out CV inside the training blocks."],
            [20, 8, 10, 40, 12, 12, 12, 12, 14, 14, 14, 14, 12, 14, 14])

    # ------------------------------------------------------------- flux baselines
    T = pd.read_csv(os.path.join(R, "baselines", "baseline_comparison_per_ion.csv")); rel = T.metabocast_within >= WITHIN_MIN
    summ = {}
    for tool in ("scFEA", "MGFEA"):
        t = T.dropna(subset=[f"{tool}_r"]); dl = t.metabocast_loso - t[f"{tool}_r"]; dw = t.metabocast_within - t[f"{tool}_r"]
        summ[tool] = dict(n=len(t), tool_median=float(t[f"{tool}_r"].median()), loso_median=float(t.metabocast_loso.median()),
                          within_median=float(t.metabocast_within.median()), loso_wins=int((dl > 0).sum()), p_loso=wp(dl),
                          within_wins=int((dw > 0).sum()), p_within=wp(dw), n_abs_gt_01=int((t[f"{tool}_r"].abs() > 0.1).sum()),
                          loso_median_diff=boot_median(dl))
    tr = T[rel]
    summ["within_supported"] = dict(n=int(rel.sum()), loso_median=float(tr.metabocast_loso.median()), scFEA_median=float(tr.scFEA_r.median()),
                                    wins=int((tr.metabocast_loso > tr.scFEA_r).sum()), p=wp(tr.metabocast_loso - tr.scFEA_r),
                                    examples={n_: (float(tr[tr.annotation == n_].metabocast_loso.iloc[0]), float(tr[tr.annotation == n_].scFEA_r.iloc[0]))
                                              for n_ in ["Glutathione", "AMP", "Xanthine"] if (tr.annotation == n_).any()})
    N["flux_baselines"] = summ
    cols = ["mz", "annotation", "adduct", "cls", "compound", "metabocast_within", "metabocast_loso", "scFEA_r", "scFEA_n_sections", "MGFEA_r", "MGFEA_n_sections"]
    fs, fr = pd.read_csv(os.path.join(R, "flux", "flux_stroma.csv")), pd.read_csv(os.path.join(R, "flux", "flux_rest.csv"))
    frows = T[cols].values.tolist()
    frows += [[], ["Warburg substrate–product test, tumour vs stroma", "tumour spots", "stroma spots", "",
                   "ROC-AUC, measured metabolites", "ROC-AUC, predicted metabolites", "ROC-AUC, enzyme expression"]]
    frows += [[x.section, int(x.n_tumour), int(x.n_other), "", x.measured, x.predicted, x.enzyme] for _, x in fs.iterrows()]
    frows.append([f"mean of {len(fs)} breast sections", "", "", "", fs.measured.mean(), fs.predicted.mean(), fs.enzyme.mean()])
    frows.append([f"sensitivity: tumour vs all other spots, mean of {len(fr)} sections", "", "", "", fr.measured.mean(),
                  fr.predicted.mean(), fr.enzyme.mean()])
    N["flux_baselines"]["warburg"] = dict(stroma={k: float(fs[k].mean()) for k in ("measured", "predicted", "enzyme")}, n=len(fs),
                                          rest={k: float(fr[k].mean()) for k in ("measured", "predicted", "enzyme")})
    B.sheet("flux_baselines", ["m/z", "putative candidate", "adduct", "class", "matched compound", "MetaboCast within-section r",
                               "MetaboCast held-out-section r", "scFEA r", "scFEA n sections", "MGFEA r", "MGFEA n sections"],
            frows,
            [f"scFEA: {summ['scFEA']['n']} ions with a single negative-mode candidate matched to a scFEA compound; median r "
             f"{summ['scFEA']['tool_median']:.3f}; MetaboCast held-out-section median {summ['scFEA']['loso_median']:.3f} (higher for "
             f"{summ['scFEA']['loso_wins']}/{summ['scFEA']['n']}, Wilcoxon P = {summ['scFEA']['p_loso']:.2g}). MGFEA: median "
             f"{summ['MGFEA']['tool_median']:.3f}; MetaboCast higher for {summ['MGFEA']['loso_wins']}/{summ['MGFEA']['n']} "
             f"(P = {summ['MGFEA']['p_loso']:.2g}).",
             "scFEA (human M168 map) on all 7 sections; MGFEA in its unsupervised spatial mode on BC_823 and LC_170; neither uses MSI. "
             "Per-ion Spearman between the tool's per-spot estimate and measured log intensity, averaged over sections where the ion is "
             "detected in >20% of spots. The held-out-section MetaboCast prediction, which never sees the section's MSI, is the matched "
             "comparison; within-section values use the section's own MSI (reproduce/baselines/compare_baselines.py).",
             "Warburg substrate–product test (reproduce/flux_validation.py; lower block): z(lactate [M−H]−, m/z 89.025) − z(hexose [M−H]−, "
             "m/z 179.055; any C6H12O6 isomer) as a one-dimensional score separating tumour from stroma spots (five-fold cross-validated "
             "ROC-AUC, logistic regression), computed from measured metabolites, from held-out-section MetaboCast predictions and, for "
             "the enzyme column, from glycolysis and TCA-cycle gene-set activity; tumour against stroma spots in the four breast "
             "sections (the lung sections carry no stroma annotation); the sensitivity row contrasts tumour with all other spots in all "
             "seven sections."],
            [22, 18, 10, 16, 16, 18, 18, 10, 10, 10, 10])

    # ------------------------------------------------------------- per-ion accuracy, tier and candidates
    W = np.load(os.path.join(R, "within_desium_per_ion.npz")); L = np.load(os.path.join(R, "loso_desium_per_ion.npz"))
    MZ = W[f"mz__{DESIUM[0]}"]
    with np.errstate(all="ignore"):
        Mw = np.vstack([W[f"MetaboCast__{s}"] for s in DESIUM])
        r, sd, n = np.nanmean(Mw, 0), np.nanstd(Mw, 0), np.sum(np.isfinite(Mw), 0)
        rt = np.nanmean(np.vstack([W[f"MetaboCast [target scale]__{s}"] for s in DESIUM]), 0)
        lr = np.nanmean(np.vstack([L[s] for s in DESIUM]), 0)
    a = pd.read_csv(os.path.join(R, "desium_ion_annotation.csv"))
    has = a.n_candidates > 0
    tiers = np.array([tier(x, y) for x, y in zip(r, lr)])
    first = a.annotation.fillna("").astype(str).str.split(";").str[0].str.strip()
    full = pd.DataFrame({"mz": np.round(MZ, 4), "n candidates": a.n_candidates,
                         "candidates (observed − theoretical m/z, mDa)": a.candidates.where(has, ""),
                         "nearest candidate": first.where(has, ""), "adduct": a.adduct.where(has, ""),
                         "class (single-class channels)": a["class"].where(has & (a["class"] != "ambiguous"), ""),
                         "within-section r (mean)": r, "s.d.": sd, "n sections": n, "within-section r (target scale)": rt,
                         "held-out-section r (mean)": lr, "tier": tiers})
    sel = full[has.values & (a["class"] != "ambiguous").values].dropna(subset=["within-section r (mean)"])
    sel = sel[sel["n sections"] >= 2].sort_values("within-section r (mean)", ascending=False)
    top = pd.concat([sel.head(12), sel.tail(8)])
    cnt = pd.Series(tiers).value_counts().to_dict()
    N["per_ion"] = dict(tiers=cnt, n_within_ge_030=int(np.nansum(r >= WITHIN_MIN)), n_transfer_ge_020=int(np.nansum(lr >= TRANSFER_MIN)),
                        n_both=int(np.nansum((r >= WITHIN_MIN) & (lr >= TRANSFER_MIN))), median_within=float(np.nanmedian(r)),
                        median_transfer=float(np.nanmedian(lr)), n_candidates=int(has.sum()), n_multi=int((a.n_candidates > 1).sum()),
                        n_single_class=int((has & (a["class"] != "ambiguous")).sum()),
                        tiers_annotated=pd.Series(tiers[has.values]).value_counts().to_dict())
    B.sheet("per_ion", ["m/z", "nearest candidate", "adduct", "class", "within-section r (mean)", "s.d. across sections", "n sections",
                        "held-out-section r (mean)", "tier"],
            top[["mz", "nearest candidate", "adduct", "class (single-class channels)", "within-section r (mean)", "s.d.", "n sections",
                 "held-out-section r (mean)", "tier"]].values.tolist(),
            ["Best 12 and worst 8 channels among those with a single-class putative candidate, detected in at least two sections. "
             "Within-section r: spatial-block CV, intensity scale, mean over sections; raised by spatial autocorrelation and not a "
             "transfer estimate. Held-out-section r: leave-one-section-out Spearman, mean over the seven held-out sections.",
             f"Tiers (MetaboCast.reliability_table): transfer-supported, held-out-section r ≥ {TRANSFER_MIN}; within-section only, "
             f"within-section r ≥ {WITHIN_MIN} with weaker transfer; low otherwise. Counts over all 2,086 channels: "
             + ", ".join(f"{k} {v}" for k, v in cnt.items()) + ".",
             "Identities are putative exact-mass candidates (negative-mode adducts, ±0.01 Da; the reference of metabocast/annotation.py, "
             "also used for the HCC/GBM applications), not confirmed identifications. Sheet "
             f"{table_sheet('per_ion')}b lists all 2,086 channels with every candidate, nearest first, and its mass error (observed "
             "minus theoretical m/z); the adduct column refers to the nearest candidate."],
            [10, 22, 10, 16, 12, 12, 10, 14, 18])
    s4b = table_sheet("per_ion") + "b"
    ws = B.wb.create_sheet(s4b); ws.append([f"Table {s4b} | All 2,086 DESIUM m/z channels: accuracy, tier and putative candidates"])
    ws["A1"].font = Font(bold=True, size=11); ws.append([]); ws.append(list(full.columns))
    for c in ws[3]:
        c.font = Font(bold=True); c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in full.itertuples(index=False):
        ws.append([clean(x) for x in row])
    ws.freeze_panes = "A4"
    for i, wdt in enumerate([10, 8, 40, 22, 10, 18, 12, 10, 8, 12, 12, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt

    # ------------------------------------------------------------- coverage of the released widths
    cov = J("coverage", "coverage_outer.json")
    rows = []
    for s in DESIUM:
        p = cov["per_section"][s]; sp = p["coverage_split_by_calibration_patient"]
        rows.append([s, p["patient"], p["n_spots"], p["n_detected"], p["coverage_crossfit"], p["coverage_split_mean"],
                     f"{min(sp.values()):.3f}-{max(sp.values()):.3f}", p["median_width"], p["median_relative_width"],
                     wj[s]["coverage90"], p["lopo_median_r"]])
    rows.append(["mean over sections", "", "", "", cov["mean_coverage_crossfit_sections"], cov["mean_coverage_split_sections"], "", "", "",
                 float(np.mean([wj[s]["coverage90"] for s in DESIUM])), cov["lopo_mean_of_sections"]])
    rows.append(["mean over patients", "", "", "", cov["mean_coverage_crossfit_patients"], cov["mean_coverage_split_patients"], "", "", "",
                 "", ""])
    N["coverage"] = dict(crossfit_sections=cov["mean_coverage_crossfit_sections"], split_sections=cov["mean_coverage_split_sections"],
                         crossfit_patients=cov["mean_coverage_crossfit_patients"], split_patients=cov["mean_coverage_split_patients"],
                         within=float(np.mean([wj[s]["coverage90"] for s in DESIUM])),
                         per_section={s: cov["per_section"][s]["coverage_crossfit"] for s in DESIUM},
                         lopo_mean=cov["lopo_mean_of_sections"], lopo_sd=cov["lopo_sd_of_sections"])
    B.sheet("coverage", ["Held-out section", "Patient", "Spots", "Ions detected", "Coverage, cross-fitted widths (released procedure)",
                         "Coverage, split calibration (mean)", "Split calibration range over calibration patients", "Median width (log1p)",
                         "Median relative width (width / s.d.)", "Within-section coverage (calibration split)", "Held-out-patient median r"],
            rows,
            ["Outer test (reproduce/coverage_outer.py): each patient is held out of fitting, penalty selection and width estimation "
             "(the two BC_515 sections together). Cross-fitted widths: every remaining patient is predicted by a model fitted on the "
             "others and the per-ion 90% quantile of the pooled absolute residuals is attached to the model fitted on all remaining "
             "patients (the procedure used for the released model, one level down). Split calibration: a model fitted on all but one "
             "remaining patient, calibrated on that patient alone.",
             "Coverage: fraction of (spot, ion) pairs of the held-out section with absolute log1p error within the width, over ions "
             "detected in >20% of its spots; nominal 0.90. Widths are empirical; no finite-sample guarantee is claimed.",
             "Within-section coverage: intervals of Fig. 4a-c (each fold model refitted on a random 75% of the spots in its training "
             "blocks, widths from the other 25% of those spots)."],
            [20, 10, 8, 10, 16, 16, 18, 12, 14, 16, 14])

    # ------------------------------------------------------------- transfer (DESIUM)
    v, vr = med("loso_desium.json"), med("loso_desium_ridge.json")
    lopo = {s: cov["per_section"][s]["lopo_median_r"] for s in DESIUM}
    base = J("loso_baselines.json")["per_section"] if os.path.exists(os.path.join(R, "loso_baselines.json")) else {}
    bm = ["graph-free", "KEGG features", "joint NMF", "single gene"]
    rows = [[s, c, v[s], vr[s], lopo[s]] + [base.get(s, {}).get(m, np.nan) for m in bm] for s, c in zip(DESIUM, CANCER)]
    allc = {"MetaboCast": v, "ridge only": vr, "patient held out": lopo, **{m: {s: base.get(s, {}).get(m, np.nan) for s in DESIUM} for m in bm}}
    for lab, idx in [("Breast mean (n=4)", slice(0, 4)), ("Lung mean (n=3)", slice(4, 7)), ("Overall mean ± s.d. (n=7)", slice(0, 7))]:
        rows.append([lab, ""] + [msd(np.array([d[s] for s in DESIUM])[idx]) for d in allc.values()])
    eff = []; N["transfer"] = dict(mean=float(np.mean(list(v.values()))), sd=float(np.std(list(v.values()), ddof=1)),
                                   breast=float(np.mean([v[s] for s in DESIUM[:4]])), lung=float(np.mean([v[s] for s in DESIUM[4:]])),
                                   ridge=float(np.mean(list(vr.values()))), lopo=float(np.mean(list(lopo.values()))))
    for name, d in [("ridge only", vr)] + [(m, allc[m]) for m in bm]:
        x = np.array([v[s] for s in DESIUM]); y = np.array([d[s] for s in DESIUM], float)
        if np.isnan(y).any():
            continue
        m_, lo_, hi_ = boot_ci(x, y); dd = x - y
        shown = {"graph-free": "no diffusion", "KEGG features": "KEGG features + regressor", "single gene": "single-gene OLS"}.get(name, name)
        eff.append([f"MetaboCast − {shown}", f"{m_:+.4f}", f"[{lo_:+.4f}, {hi_:+.4f}]", f"{int((dd > 0).sum())}/7", round(wp(dd), 4)])
        N["transfer"][f"vs {name}"] = dict(delta=m_, ci=(lo_, hi_), wins=int((dd > 0).sum()), p=wp(dd), other_mean=float(y.mean()))
    rows += [[], ["Paired comparison over the 7 held-out sections", "Δ", "95% CI (paired bootstrap)", "sections improved", "Wilcoxon P"]] + eff
    B.sheet("transfer", ["Held-out section", "Cancer", "MetaboCast", "MetaboCast, ridge only", "MetaboCast, patient held out"]
            + ["no diffusion", "KEGG features + regressor", "joint NMF", "single-gene OLS"], rows,
            ["Detection-matched median Spearman on the held-out section (ions detected in >20% of its spots). Each model is fitted on "
             "the six other sections (highly variable genes and PCA from the training sections; reproduce/loso_benchmark.py). Patient "
             "held out: both BC_515 sections leave training together when either is tested (reproduce/coverage_outer.py).",
             "Paired-supervision baselines under the same held-out design (reproduce/loso_baselines.py): the released estimator without "
             "diffusion features (no diffusion) or on KEGG pathway features with one diffused copy; joint NMF (20 factors) of training genes and "
             "metabolites, with held-out metabolites reconstructed from held-out genes; signed regression on the single training gene "
             "most correlated with each ion."], [26, 14, 12, 14, 16, 14, 14, 12, 12])

    # ------------------------------------------------------------- histology
    hl = A("histology_loso.csv").pivot(index="sample", columns="model", values="median_spearman")
    ho = A("histology_only_desium.csv"); hw = A("histology_within.csv").pivot(index="sample", columns="model", values="median_spearman")
    he = ho[ho.model == "H&E-only"].pivot(index="section", columns="test", values="median_spearman")
    sh = ho[ho.model == "H&E-only(shuffled)"].set_index("section").median_spearman
    hs = [s for s in DESIUM if s in hl.index]
    rows = [[s, hl.loc[s, "MetaboCast"], he.loc[s, "LOSO (cross-section)"], hl.loc[s, "MetaboCast+H&E"], sh.get(s, np.nan),
             hw.loc[s, "MetaboCast"], he.loc[s, "within-sample 5-fold"], hw.loc[s, "MetaboCast+H&E"]] for s in hs]
    mrow = lambda col: float(np.mean([x[col] for x in rows]))
    rows.append([f"mean ({len(hs)} sections)"] + [mrow(i) for i in range(1, 8)])
    d_he = np.array([he.loc[s, "LOSO (cross-section)"] - hl.loc[s, "MetaboCast"] for s in hs])
    d_both = np.array([hl.loc[s, "MetaboCast+H&E"] - hl.loc[s, "MetaboCast"] for s in hs])
    tx_ = [hl.loc[s, "MetaboCast"] for s in hs]
    N["histology"] = dict(transcriptome=mrow(1), he_only=mrow(2), combined=mrow(3), shuffled=mrow(4), within_transcriptome=mrow(5),
                          within_he=mrow(6), within_combined=mrow(7), he_wins=int((d_he > 0).sum()), combined_wins=int((d_both > 0).sum()),
                          n=len(hs), p_he=wp(d_he), p_combined=wp(d_both),
                          ci_he=boot_ci([he.loc[s, "LOSO (cross-section)"] for s in hs], tx_)[1:],
                          ci_combined=boot_ci([hl.loc[s, "MetaboCast+H&E"] for s in hs], tx_)[1:])
    B.sheet("histology", ["Held-out section", "Transcriptome only (MetaboCast)", "H&E only", "Transcriptome + H&E",
                          "H&E only, shuffled features", "Within: MetaboCast", "Within: H&E only", "Within: MetaboCast + H&E"], rows,
            ["Median Spearman over the 1,000 most variable ions of each held-out section; held-out-section design (columns 2-5) and "
             "five spatial blocks within each section (columns 6-8). H&E features: 13 per-spot patch statistics from the embedded Visium "
             "images (RGB means and s.d., optical-density means, grayscale mean and s.d., an edge-gradient summary, tissue fraction), "
             "appended via extra_key or used alone (metabocast.MorphologyMetabolitePredictor). Transcriptome arms use n_hvg = 2,000.",
             "BC_525 is excluded: its embedded image could not be matched unambiguously to its Visium library. Shuffling the H&E features "
             "across spots removes the signal. The shipped model is transcriptome-only.",
             f"H&E only − transcriptome only: higher in {int((d_he > 0).sum())}/{len(hs)} sections (Wilcoxon P = {wp(d_he):.2f}); "
             f"transcriptome + H&E − transcriptome only: higher in {int((d_both > 0).sum())}/{len(hs)} (P = {wp(d_both):.3f})."],
            [24, 16, 12, 16, 16, 14, 14, 16])

    # ------------------------------------------------------------- SMA registration and cohorts
    au = pd.read_csv(os.path.join(R, "sma", "sma_registration_audit.csv")).set_index("section")
    qc = pd.read_csv(os.path.join(R, "sma_coregistration_qc.csv")).set_index("section")
    rows = []
    for s in au.index:
        x, q = au.loc[s], qc.loc[s]
        rows.append([s.replace("SMA_", ""), x.matrix, q.panel, q.species, x.subject, q.region, q.spots_in_tissue, x.published_pairs,
                     q.spots, q.n_mz, x.geometry_orientation, x.geometry_iou, x.geometry_iou_runner_up, x.geometry_within_1px,
                     x.get("coupling_within_1px", np.nan), q.within_section_coupling, q.shuffled_null])
    N["sma_registration"] = dict(geometry_within_1px=au.geometry_within_1px.round(3).to_dict(),
                                 coupling_within_1px=au.coupling_within_1px.round(3).to_dict() if "coupling_within_1px" in au else {},
                                 spots={s: int(qc.loc[s, "spots"]) for s in au.index},
                                 coupling=qc.within_section_coupling.round(3).to_dict(), null_max=float(qc.shuffled_null.abs().max()))
    B.sheet("sma_registration", ["Section", "Matrix", "Panel", "Species", "Subject", "Region", "In-tissue spots", "Published pairs",
                                 "Spots used", "m/z channels", "Geometry: orientation", "Geometry: IoU", "Geometry: IoU runner-up",
                                 "Geometry: within 1 px of published", "Coupling-chosen orientation: within 1 px of published",
                                 "Coupling after registration", "Shuffled-spot null"], rows,
            ["Every SMA analysis uses the spot-to-pixel pairing released by the SMA authors (deposit misc.zip: knn_spatial_df_filtered_list "
             "for the mouse sections; hPDStr.multi.list for the human sections, matched to raster pixels by identical spectra). The authors "
             "aligned the images by hand in STUtility and paired each spot with its nearest pixel (pairs more than 35 image pixels apart "
             "dropped; one pixel per spot). The pairing was fixed before any prediction here (pipeline/sma_published_pairs.R/.py, "
             "pipeline/sma_to_metabocast.py). Spots without a pair are not used.",
             "Subject: animal or donor from the deposit's metadata.csv. The DHB and 9-AA sections come from one mouse (mPD3); the six FMP "
             "mouse sections come from three animals (mPD1, mPD3, mPD4), each with a striatum and a substantia nigra section; the four "
             "human sections come from one donor.",
             "Alternatives compared with the published pairing (reproduce/sma_registration_audit.py): tissue-geometry registration, "
             "which does not use the outcome (orientation, scale and shift maximising the overlap of the MALDI tissue mask with Space "
             "Ranger's in-tissue spots), and the orientation that maximises transcriptome-metabolite coupling, which does. The human "
             "sections fill the capture area, so their outline does not fix the orientation; coupling-chosen orientations disagree with "
             "the published pairing for most sections, so coupling is not a valid registration criterion.",
             "Coupling after registration: median per-ion Spearman of a 30-PC ridge on random halves of the section, with a shuffled-spot "
             "null; a quality check computed after registration that plays no part in it."],
            [16, 8, 14, 8, 8, 14, 10, 10, 10, 10, 14, 10, 12, 14, 18, 12, 12])

    # ------------------------------------------------------------- SMA within-section benchmark
    sj = J("within_sma.json")
    groups = [("DESIUM (DESI, human tumour)", DESIUM, wj), ("SMA DHB lipids (mouse striatum)", SMA_GROUPS["dhb"], sj),
              ("SMA 9-AA metabolites (mouse striatum)", ["SMA_9AA_V11L12-038_D1"], sj),
              ("SMA FMP mouse striatum", [s for s in SMA_GROUPS["fmp_mouse"] if s not in NIGRA], sj),
              ("SMA FMP mouse substantia nigra", [s for s in SMA_GROUPS["fmp_mouse"] if s in NIGRA], sj),
              ("SMA FMP human striatum", SMA_GROUPS["fmp_human"], sj)]
    meth7 = ["MetaboCast", "MetaboCast + KEGG", "MetaboCast (ridge only)", KEGG, "joint-NMF", "no diffusion", "single-gene OLS"]
    rows = []; N["sma_within"] = {}
    for lab_, ss, js_ in groups:
        ss = [s for s in ss if s in js_]
        rows.append([lab_, len(ss)] + [msd([js_[s]["scores"].get(m, np.nan) for s in ss]) for m in meth7]
                    + [msd([js_[s]["full_median"] for s in ss])])
        N["sma_within"][lab_] = {m: [js_[s]["scores"].get(m, np.nan) for s in ss] for m in meth7}
        N["sma_within"][lab_]["all_detected"] = [js_[s]["full_median"] for s in ss]
    rows += [[], ["Per section", ""]]
    for s in [s for g in SMA_GROUPS.values() for s in g] + ["SMA_9AA_V11L12-038_D1"]:
        if s in sj:
            rows.append([s.replace("SMA_", ""), sj[s]["n_panel"]] + [sj[s]["scores"].get(m, np.nan) for m in meth7] + [sj[s]["full_median"]])
    fm = [s for s in SMA_GROUPS["fmp_mouse"] + SMA_GROUPS["fmp_human"] if s in sj]
    dn = np.array([sj[s]["scores"]["MetaboCast"] - sj[s]["scores"]["joint-NMF"] for s in fm])
    dg = np.array([sj[s]["scores"]["MetaboCast"] - sj[s]["scores"]["no diffusion"] for s in fm])
    allsma = [s for s in sj]
    da = np.array([sj[s]["scores"]["MetaboCast"] - sj[s]["scores"]["joint-NMF"] for s in allsma])
    ci0 = lambda d_: boot_ci(d_, np.zeros_like(d_))[1:]
    N["sma_within_tests"] = dict(fmp_vs_nmf=dict(delta=float(dn.mean()), wins=int((dn > 0).sum()), p=wp(dn), n=len(dn), ci=ci0(dn)),
                                 fmp_vs_graphfree=dict(delta=float(dg.mean()), wins=int((dg > 0).sum()), p=wp(dg), n=len(dg), ci=ci0(dg)),
                                 all_vs_nmf=dict(delta=float(da.mean()), wins=int((da > 0).sum()), p=wp(da), n=len(da), ci=ci0(da)),
                                 diffusion_gain={lab_: float(np.mean([js_[s]["scores"]["MetaboCast"] - js_[s]["scores"]["no diffusion"]
                                                                      for s in ss if s in js_])) for lab_, ss, js_ in groups})
    B.sheet("sma_within", ["Cohort", "n sections (per-section rows: panel size)"] + meth7 + ["MetaboCast, all detected ions"], rows,
            [f"FMP (10 sections): MetaboCast − joint NMF = {dn.mean():+.3f}, higher in {int((dn > 0).sum())}/{len(dn)}, exact Wilcoxon "
             f"P = {wp(dn):.4f}; MetaboCast − no diffusion = {dg.mean():+.3f}, {int((dg > 0).sum())}/{len(dg)}, P = {wp(dg):.4f}.",
             f"Same harness as {TR('benchmark')} (reproduce/within_benchmark.py); median Spearman over the 1,000 most variable detected ions (all "
             "detected ions when fewer pass). The KEGG variants need human gene symbols, so they equal MetaboCast or are n/a for mouse.",
             f"Registration: the published spot-to-pixel pairing ({TR('sma_registration')})."], [36, 12, 14, 14, 14, 16, 12, 16, 12, 16])

    # ------------------------------------------------------------- SMA transfer
    S = {g: J("sma", f"loso_{g}.json")["per_section"] for g in SMA_GROUPS}
    Sr = {g: med("sma", f"loso_{g}_ridge.json") for g in SMA_GROUPS}
    LA = J("sma", "loso_fmp_mouse_by_animal.json")["per_section"]
    rows = []
    for g, ss in SMA_GROUPS.items():
        for s in ss:
            x = S[g][s]
            rows.append([s.replace("SMA_", ""), {"dhb": "DHB lipids", "fmp_mouse": "FMP mouse", "fmp_human": "FMP human"}[g],
                         x.get("subject", ""), "substantia nigra" if s in NIGRA else "striatum", x["n_detected"], x.get("n_train_sections", ""),
                         x["median"], Sr[g][s], LA[s]["median"] if s in LA else np.nan, LA[s].get("n_train_sections", "") if s in LA else ""])
    rows.append([])
    N["sma_transfer"] = {}
    for g, lab_ in [("dhb", "DHB lipids (2 sections, 1 animal)"), ("fmp_mouse", "FMP mouse (6 sections, 3 animals)"),
                    ("fmp_human", "FMP human (4 sections, 1 donor)")]:
        vals = [S[g][s]["median"] for s in SMA_GROUPS[g]]
        rows.append([lab_, "mean", "", "", "", "", np.mean(vals), np.mean([Sr[g][s] for s in SMA_GROUPS[g]]),
                     np.mean([LA[s]["median"] for s in SMA_GROUPS[g]]) if g == "fmp_mouse" else np.nan, ""])
        N["sma_transfer"][g] = dict(mean=float(np.mean(vals)), per={s: S[g][s]["median"] for s in SMA_GROUPS[g]},
                                    ridge=float(np.mean([Sr[g][s] for s in SMA_GROUPS[g]])))
    st = [S["fmp_mouse"][s]["median"] for s in SMA_GROUPS["fmp_mouse"] if s not in NIGRA]
    ni = [S["fmp_mouse"][s]["median"] for s in SMA_GROUPS["fmp_mouse"] if s in NIGRA]
    N["sma_transfer"]["fmp_mouse_striatum"] = float(np.mean(st)); N["sma_transfer"]["fmp_mouse_nigra"] = float(np.mean(ni))
    N["sma_transfer"]["fmp_mouse_by_animal"] = dict(mean=float(np.mean([LA[s]["median"] for s in LA])),
                                                    striatum=float(np.mean([LA[s]["median"] for s in LA if s not in NIGRA])),
                                                    nigra=float(np.mean([LA[s]["median"] for s in LA if s in NIGRA])),
                                                    per={s: LA[s]["median"] for s in LA})
    B.sheet("sma_transfer", ["Held-out section", "Cohort", "Subject", "Region", "Ions detected (>20%)", "Training sections",
                             "MetaboCast, section held out", "Ridge only, section held out", "MetaboCast, animal held out",
                             "Training sections (animal held out)"], rows,
            ["Detection-matched median Spearman on the held-out section with the released package (3,000 training-section highly variable "
             "genes; reproduce/loso_benchmark.py). Animal held out (--hold-out subject): every section of the held-out section's animal "
             "leaves training, so the model learns from the other two animals only.",
             "DHB: the model for each section is trained on the other section of the same animal (same slide), so this is within-animal "
             "transfer between sections. Human FMP: four sections from one donor. 9-AA has a single section and no transfer estimate.",
             f"Registration: the published spot-to-pixel pairing ({TR('sma_registration')})."], [20, 12, 8, 16, 10, 10, 14, 14, 14, 14])

    # ------------------------------------------------------------- HCC-2 normal -> PVTT
    s8 = X("S8_supported_annotated.csv")
    s8 = s8.assign(_o=s8.tier.map({"transfer-supported": 0, "within-section only": 1})).sort_values(["_o", "delta_NtoP_log"],
                                                                                                      ascending=[True, False])
    B.sheet("hcc_pvtt", ["m/z", "putative candidate(s)", "DESIUM within-section r", "DESIUM held-out-section r", "tier",
                         "mean normal (s.d.)", "mean PVTT (s.d.)", "Δ normal→PVTT (predicted log)", "more than one candidate"],
            [[x.mz, x.annotation, x.within_section_r, x.transfer_r, x.tier, f"{x.mean_N:.3f} ({x.sd_N:.3f})", f"{x.mean_P:.3f} ({x.sd_P:.3f})",
              x.delta_NtoP_log, bool(x.ambiguous_annotation)] for x in s8.itertuples()],
            ["Predictions only: HCC has no measured metabolome. One cross-cancer run: MetaboCast fitted on all seven DESIUM sections, genes "
             "shared by DESIUM, the five HCC samples and both GBM sections, 100 PCs, k = 6 (reproduce/crosscancer/regenerate_crosscancer.py). "
             f"HCC-2 normal region {int(X('S8_full_all_ions.csv').n_spots_N.iloc[0])} spots, PVTT {int(X('S8_full_all_ions.csv').n_spots_P.iloc[0])}.",
             f"Rows: annotated ions in the two upper DESIUM accuracy tiers ({TR('per_ion')}b), transfer-supported first. Spots within a region are "
             "spatially dependent and come from one patient, so no per-spot test is reported; the s.d. describes spread across spots. "
             "Identities are putative exact-mass candidates of a curated list of tissue metabolites (negative-mode adducts, ±0.01 Da)."],
            [10, 46, 14, 14, 18, 16, 16, 16, 12])
    N["hcc_pvtt"] = dict(n=len(s8), by_tier=s8.tier.value_counts().to_dict())

    # ------------------------------------------------------------- HCC cross-patient contrasts
    s9 = X("S9_and_Fig7c_source.csv"); exp9 = dict(EXPECTED, Glutamate=1)
    rows = []
    lp = X("HCC_NtoT_annotated_by_patient.csv")             # full-precision contrasts, rounded once below
    full_d = {(round(float(m_), 4), p_): float(v_) for m_, p_, v_ in zip(lp.mz, lp.patient, lp.delta_NtoT)}
    s8x = X("S8_full_all_ions.csv"); dNP = dict(zip(np.round(s8x.mz.astype(float), 4), s8x.delta_NtoP_log.astype(float)))
    for _, x in s9.iterrows():
        k_ = round(float(x.mz), 4)
        vv = np.array([full_d[(k_, p_)] for p_ in PATIENTS]); e = exp9.get(x.metabolite)
        rows.append([x.metabolite, x.mz, x.within_section_r, x.transfer_r, x.tier, bool(x.ambiguous_annotation), *np.round(vv, 4),
                     round(dNP[k_], 4), int((vv > 0).sum()), ("up" if e == 1 else "down") if e else "—",
                     f"{int((np.sign(vv) == e).sum())}/4" if e else "—"])
    wide = lp.pivot_table(index=["mz", "annotation", "tier"], columns="patient", values="delta_NtoT").reset_index()
    rows += [[], ["All annotated ions in the two upper tiers", "m/z", "", "", "tier", "", "HCC-1 N→T", "HCC-2 N→T", "HCC-3 N→T", "HCC-4 N→T"]]
    rows += [[x.annotation, x.mz, "", "", x.tier, ""] + [round(float(x[p]), 4) for p in PATIENTS] for _, x in wide.iterrows()]
    B.sheet("hcc_patients", ["metabolite", "m/z", "DESIUM within-section r", "DESIUM held-out-section r", "tier", "more than one candidate",
                             "HCC-1 N→T", "HCC-2 N→T", "HCC-3 N→T", "HCC-4 N→T", "HCC-2 N→PVTT", "patients with tumour > normal",
                             "literature-expected direction", "patients matching"], rows,
            ["Difference in mean predicted log intensity, tumour minus normal (sample.ident region labels); predictions only. Each patient "
             "contributes one contrast; the four contrasts are not replicates of one experiment and no test is applied. HCC-5 has no "
             "normal region.", "The curated ions (upper block) are their [M−H]− channels; the lower block lists every annotated ion of "
             "the two upper DESIUM tiers with its candidates, nearest first.",
             f"Same run as {TR('hcc_pvtt')}, {TR('gbm_protein')} and Fig. 7."], [24, 10, 14, 14, 18, 12, 11, 11, 11, 11, 13, 14, 14, 12])

    # sheet b: the same contrasts with the application model refitted at the optimum of an extended penalty grid
    ps_ = J("crosscancer", "penalty_sensitivity.json"); pc_ = X("penalty_sensitivity_contrasts.csv")
    pc_["annotation"] = pc_["annotation"].fillna("")
    a_rel, a_ext = ps_["released_alpha"], ps_["extended_alpha"]
    clab = {"pvtt": "HCC-2 normal → PVTT", "region": "tumour − normal region", "type": "tumour- − normal-typed spots, pooled",
            "edge": "tumour- − normal-typed spots, leading-edge section"}
    rows = []
    for c_ in ("pvtt", "region", "type", "edge"):
        for set_, lab2 in (("annotated_upper_tiers", "annotated, two upper tiers"), ("transfer_supported", "all transfer-supported")):
            g_ = ps_["agreement"][c_][set_]
            rows.append([clab[c_], lab2, g_["n"], g_["spearman"], g_["same_sign"], g_["median_magnitude_ratio"]])
    rows += [[], ["Inner CV score (median per-ion Pearson, leave-one-section-out)", f"{pw(a_rel)}: {ps_['inner_cv_at_released']:.4f}",
                  f"{pw(a_ext)}: {ps_['inner_cv_at_extended']:.4f}"],
             ["Mean per-ion spatial s.d. of the HCC predictions (log units)", f"{pw(a_rel)}: {ps_['mean_spatial_sd']['released']:.4f}",
              f"{pw(a_ext)}: {ps_['mean_spatial_sd']['extended']:.4f}"]]
    for k_, a_ in (("released", a_rel), ("extended", a_ext)):
        ct_ = ps_["consistent_tumour_normal"][k_]
        rows.append([f"Annotated transfer-supported channels with the same tumour − normal direction in all four patients, {pw(a_)}",
                     "lower: " + "; ".join(ct_["lower_in_all_four"]), "higher: " + "; ".join(ct_["higher_in_all_four"])])
    PAT_ = ["HCC-1", "HCC-2", "HCC-3", "HCC-4"]
    rows += [[], ["m/z", "putative candidates", "tier", "contrast"] + [f"{p} ({pw(a_rel)})" for p in PAT_]
             + [f"{p} ({pw(a_ext)})" for p in PAT_] + ["same sign"]]
    ann_ = pc_[pc_.annotation != ""]
    order_ = ann_[["mz", "annotation", "tier"]].drop_duplicates()
    order_ = order_.assign(_t=order_.tier != "transfer-supported").sort_values(["_t", "mz"])
    for _, q_ in order_.iterrows():
        for c_ in ("pvtt", "region", "type", "edge"):
            g_ = ann_[(ann_.mz == q_.mz) & (ann_.contrast == c_)].set_index("patient")
            if not len(g_):
                continue
            v0 = [g_.delta_released.get(p, np.nan) for p in PAT_]; v1 = [g_.delta_extended.get(p, np.nan) for p in PAT_]
            same = sum(1 for x, y in zip(v0, v1) if np.isfinite(x) and np.sign(x) == np.sign(y))
            rows.append([q_.mz, q_.annotation, q_.tier, clab[c_]] + v0 + v1 + [f"{same}/{int(np.isfinite(v0).sum())}"])
    # the refit at the released penalty against the published run, wherever the published values are stored
    pub_ = []
    s8_ = X("S8_full_all_ions.csv"); s9m_ = X("S9_and_Fig7c_source.csv").set_index("metabolite")
    nt_ = X("HCC_NtoT_annotated_by_patient.csv")
    near = lambda a_, b_: np.isclose(a_, b_, atol=5e-4)
    for _, q_ in pc_[pc_.annotation != ""].iterrows():
        if q_.contrast == "pvtt":
            hit = s8_[near(s8_.mz, q_.mz)]
            pub_ += [(q_.delta_released, float(hit.delta_NtoP_log.iloc[0]))] if len(hit) else []
        elif q_.contrast == "region":
            hit = nt_[(nt_.patient == q_.patient) & near(nt_.mz, q_.mz)]
            pub_ += [(q_.delta_released, float(hit.delta_NtoT.iloc[0]))] if len(hit) else []
    for c_, fn_ in (("type", "HCC_TypeLevel_TumorCell_vs_NormalCell.csv"), ("edge", "HCC_TypeLevel_LeadingEdge.csv")):
        for _, r_ in X(fn_).iterrows():
            hit = pc_[(pc_.contrast == c_) & (pc_.patient == r_.patient) & near(pc_.mz, float(s9m_.loc[r_.metabolite, "mz"]))]
            pub_ += [(float(hit.delta_released.iloc[0]), float(r_.delta))] if len(hit) else []
    pub_ = np.array(pub_)
    agree_pub = dict(n=int(len(pub_)), max_abs_diff=float(np.abs(pub_[:, 0] - pub_[:, 1]).max()),
                     sign_changes=int((np.sign(pub_[:, 0]) != np.sign(pub_[:, 1])).sum()),
                     n_values=int(sum(len(v) for v in [pc_[pc_.annotation != ""]])))
    N["hcc_penalty"] = {k: ps_[k] for k in ("released_alpha", "extended_alpha", "inner_cv_at_released", "inner_cv_at_extended",
                                            "mean_spatial_sd", "agreement", "consistent_tumour_normal")}
    N["hcc_penalty"]["refit_vs_published"] = agree_pub
    B.sheet_b("hcc_patients", "Ridge-penalty sensitivity of the HCC contrasts",
              ["Contrast", "Channels", "n (channel × patient)", f"Spearman, {pw(a_rel)} vs {pw(a_ext)}", "same sign",
               f"median |Δ at {pw(a_ext)}| / |Δ at {pw(a_rel)}|"], rows,
              [f"The application model (all seven DESIUM sections, genes shared with HCC and GBM) selects {pw(a_rel)}, the largest value "
               f"of the default penalty grid; on a grid extended to 10¹² the inner cross-validation score peaks at {pw(a_ext)} "
               "(reproduce/crosscancer/penalty_sensitivity.py). Both models predict the four HCC patients with normal and tumour regions; "
               "differences are in mean predicted log intensity, as in the main sheet. Predictions only; no metabolites were measured.",
               "The score that selects the penalty does not depend on the scale of the predictions, so the larger penalty mainly shrinks "
               "their amplitude; the magnitudes of the differences are therefore model-dependent. Differences between separately acquired "
               "sections (region contrasts) were the least stable to the change, and differences within the leading-edge section the most.",
               "Lower block: every annotated channel of the two upper DESIUM tiers; 'same sign' counts patients whose difference has the "
               "same sign at both penalties.",
               f"Both penalties were refitted here in the pinned environment (requirements-lock.txt). The published HCC run behind "
               f"{TR('hcc_pvtt')}, {TR('hcc_patients')} and Fig. 7 used newer library versions (crosscancer/provenance.json), so its "
               "values at the released penalty differ slightly from the 10⁷ columns: for the "
               f"{agree_pub['n']} of the {agree_pub['n_values']} values in the lower block whose published counterparts are stored in "
               "reproduce/results (normal-to-PVTT and tumour-minus-normal region differences of every channel, type-level differences "
               f"of the 13 curated ions), by at most {agree_pub['max_abs_diff']:.3f} log units, with {agree_pub['sign_changes']} sign "
               "changes, all in differences close to zero."],
              [44, 34, 16, 16, 12, 18, 12, 12, 12, 12, 12, 12, 10])

    # ------------------------------------------------------------- literature and proposed tests
    s9i = s9.set_index("metabolite"); s8f = X("S8_full_all_ions.csv"); pats = [f"{p} N->T" for p in PATIENTS]

    def cnt(m, sign):
        vv = s9i.loc[m, pats].values.astype(float); hit = [p.split()[0] for p, x in zip(pats, vv) if np.sign(x) == sign]
        return f"{len(hit)}/4" + (f" ({', '.join(hit)})" if hit else "")
    lac = s8f.loc[(s8f.mz - 89.0244).abs().idxmin()]
    items = [("Arachidonate ↑", "Arachidonate(FA20:4)", 1, "eicosanoid precursor; PUFA accumulation in HCC",
              "targeted MSI of FA20:4 with PLA2G4A/ALOX staining on adjacent sections"),
             ("DHA ↑", "DHA(FA22:6)", 1, "PUFA; reports conflict (loss with progression in some studies)",
              "targeted MSI of FA22:6; check against PUFA-phospholipid species"),
             ("Hexose ↓ (any C6H12O6 isomer, inositol included)", "Hexose", -1, "glycolytic consumption",
              "MSI with glucose/inositol separation (MS/MS)"),
             ("Glutathione ↑", "Glutathione(GSH)", 1, "antioxidant; GPX4-dependent ferroptosis resistance",
              "proposed: MSI or LC-MS of GSH/GSSG in PVTT versus adjacent normal, with GPX4 and SLC7A11 staining and RSL3 or erastin "
              "sensitivity of PVTT-derived cells"),
             ("Ascorbate ↑", "Ascorbate", 1, "co-antioxidant", "targeted measurement; DESIUM transfer near zero"),
             ("Glutamate ↑", "Glutamate", 1, "glutaminolysis feeding GSH synthesis", "isotope tracing; DESIUM transfer near zero"),
             ("Palmitate ↓", "Palmitate(FA16:0)", -1, "saturated free fatty acid depletion / lipid remodelling", "DESIUM transfer near zero: not interpretable"),
             ("Stearate ↓", "Stearate(FA18:0)", -1, "saturated free fatty acid depletion / lipid remodelling", "DESIUM transfer near zero: not interpretable")]
    def channel(v):
        """m/z and adduct of the nearest candidate of a channel (S8_full_all_ions annotation)."""
        row = s8f.loc[(s8f.mz - v).abs().idxmin()]
        return f"{row.mz:.3f} {str(row.annotation).split(';')[0].split(' ', 1)[1].replace('-', '−')}"
    rows = [[lab_, channel(s9i.loc[m, "mz"]), bio, s9i.loc[m, "tier"], s9i.loc[m, "transfer_r"], cnt(m, sg),
             f"{s9i.loc[m, 'HCC-2 N->P']:+.3f}", test] for lab_, m, sg, bio, test in items]
    lr_ = X("crosscancer_ion_reliability.csv").set_index("metabolite").loc["Lactate"]
    lv = np.array([lr_[f"{p} tumour - normal"] for p in PATIENTS], float); lhit = [p for p, x in zip(PATIENTS, lv) if x > 0]
    rows.append(["Lactate ↑", channel(lac.mz), "elevated in HCC (Warburg)", lac.tier, lac.transfer_r,
                 f"{len(lhit)}/4" + (f" ({', '.join(lhit)})" if lhit else ""), f"{lac.delta_NtoP_log:+.3f}",
                 "not interpretable (low DESIUM accuracy); measure lactate directly, for example by targeted MSI"])
    B.sheet("hcc_literature", ["Predicted change", "Channel (m/z, adduct)", "Known HCC biology", "DESIUM tier",
                               "DESIUM held-out-section r", "Patients matching (N→T)", "HCC-2 N→PVTT", "What would test it"], rows,
            ["Predictions are hypotheses, not measurements. Transfer-supported ions are listed first; for within-section-only ions the "
             "DESIUM evidence does not carry over to unseen sections, so their HCC predictions are weak leads at most.",
             f"Per-patient values: {TR('hcc_patients')} for the curated ions and {TR('crosscancer_summary')} for lactate. Hexose stands "
             "for every C6H12O6 isomer, inositol included."],
            [26, 16, 44, 18, 14, 18, 14, 60])

    # ------------------------------------------------------------- GBM protein co-localisation
    s10 = X("S10_metabolite_x_protein.csv")
    head10 = {"predicted_metabolite": "predicted ion (putative candidate)", "mz": "m/z", "within_section_r": "DESIUM within-section r",
              "transfer_r": "DESIUM held-out-section r", "tier": "tier", "ambiguous_annotation": "more than one candidate"}
    B.sheet("gbm_protein", [head10.get(c_, c_) for c_ in s10.columns], s10.values.tolist(),
            ["Exploratory. Per-spot Spearman between predicted metabolites and 31 measured antibody markers (isotype controls removed, "
             "log1p counts) on the GBM CytAssist section. Part of the antibody signal can be reconstructed from RNA, so these are "
             "co-localisations of two transcriptome-derived or transcriptome-correlated signals, not validation of the predictions.",
             "N-acetylaspartate, a neuronal metabolite, follows proliferation markers here; the DESIUM training panel contains no neural "
             "tissue, so we read this as a possible failure of transfer to brain."])

    # ------------------------------------------------------------- HCC/GBM summary
    xc = J("crosscancer", "crosscancer_summary.json"); ion = X("crosscancer_ion_reliability.csv")
    hc = A("histology_only_hcc_concordance.csv").set_index("metabolite"); cn = xc["concordance"]
    he_k = {p: sum(int(np.sign(hc.loc[m, f"{p} H&E delta"]) == e) for m, e in EXPECTED.items()) for p in PATIENTS}
    rows = [[f"annotated ions on the 2,086-channel axis (the reference of {TR('per_ion')}b, negative-mode adducts)", xc["n_annotated"]]]
    rows += [[f"  of which {k}", v] for k, v in xc["n_annotated_by_tier"].items()]
    rows += [["median DESIUM within-section r, annotated ions", xc["median_within_r_annotated"]],
             ["median DESIUM held-out-section r, annotated ions", xc["median_loso_r_annotated"]],
             ["literature-expected normal→tumour directions matched, transcriptome predictions (7 metabolites × 4 patients)",
              f"{cn['k']}/28 (descriptive)"],
             ["  per patient (HCC-1, -2, -3, -4)", ", ".join(f"{cn['per_patient'][p]}/7" for p in PATIENTS)],
             ["the same, H&E-only model on the HCC images", f"{sum(he_k.values())}/28 (descriptive)"],
             ["  per patient (HCC-1, -2, -3, -4)", ", ".join(f"{he_k[p]}/7" for p in PATIENTS)],
             ["palmitate / stearate predicted lower in tumour (patients)",
              f"{xc['sfa_depletion_patients']['Palmitate(FA16:0)']} / {xc['sfa_depletion_patients']['Stearate(FA18:0)']}"],
             ["HCC-2 normal→PVTT literature-expected directions", f"{xc['HCC-2 N->P expected directions']}/7"],
             ["GBM CytAssist: mean of each metabolite's strongest |ρ| with a protein", xc["gbm_mean_max_abs_rho"]],
             [], ["metabolite"] + list(ion.columns[1:])]
    rows += [list(x.values) for _, x in ion.iterrows()]
    rows += [[], ["H&E-only model, normal→tumour difference per patient", "expected", "concordant"] + [f"{p} H&E Δ" for p in PATIENTS]]
    rows += [[m, hc.loc[m, "expected_N_to_T"], hc.loc[m, "H&E-only concordant"]] + [hc.loc[m, f"{p} H&E delta"] for p in PATIENTS]
             for m in EXPECTED]
    N["crosscancer_summary"] = dict(xc, he_only_k=sum(he_k.values()), he_only_per_patient=he_k)
    B.sheet("crosscancer_summary", ["measure", "value", "", "", "", "", "", ""], rows,
            ["HCC and GBM have no measured metabolome, so accuracy cannot be computed there. The table reports the measured DESIUM "
             "accuracy and tier of the ions those analyses use, and counts of agreement with literature-expected directions. The 28 "
             "direction calls share one model, seven metabolites and four patients; they are not independent replicates and no test is "
             "applied (reproduce/crosscancer/summarise_crosscancer.py; H&E-only predictions from "
             "reproduce/crosscancer/regenerate_s14_panelBC.py)."], [70, 18, 18, 18, 18, 18, 18, 18])

    # ------------------------------------------------------------- harmonisation
    h = J("harmonization.json")
    lab = {"per-fold PCA": "PCA fitted on the six training sections (default)", "pooled PCA": "PCA fitted on all seven sections (transductive)",
           "pooled + Harmony": "pooled PCA integrated across sections with Harmony", "centred PCA": "per-fold PCA, scores centred within each section",
           "z-scored features": "diffusion features z-scored within each section"}
    rows = [[lab[k], h[k]["overall"][0], h[k]["overall"][1], h[k]["breast"], h[k]["lung"]] + [h[k]["per"][s] for s in DESIUM] for k in lab]
    N["harmonisation"] = {k: h[k]["overall"] for k in lab}
    B.sheet("harmonisation", ["Transcriptome representation", "mean", "s.d.", "breast", "lung"] + DESIUM, rows,
            ["The MetaboCast estimator with only the transcriptome representation changed; 2,000 highly variable genes pooled over the "
             "seven sections; detection-matched median Spearman per held-out section (reproduce/harmonization.py). Harmony uses no "
             "metabolite information. The released package (3,000 training-section HVGs) gives the values of " + TR("transfer") + "."],
            [52, 11, 9, 9, 9] + [10] * 7)

    # ------------------------------------------------------------- KEGG prior
    kd = med("loso_desium_kegg.json"); dk = np.array([kd[s] - v[s] for s in DESIUM])
    mk, lk, hk = boot_ci([kd[s] for s in DESIUM], [v[s] for s in DESIUM])
    rows = [["DESIUM (DESI, human tumour, n=7)", np.mean(list(v.values())), np.mean(list(kd.values())), dk.mean(),
             f"95% CI [{lk:+.4f}, {hk:+.4f}]; {int((dk > 0).sum())}/7 sections improve"]]
    N["kegg_prior"] = {"DESIUM": dict(base=float(np.mean(list(v.values()))), kegg=float(np.mean(list(kd.values()))), delta=float(dk.mean()),
                                      ci=(lk, hk), wins=int((dk > 0).sum()), breast=float(dk[:4].mean()), lung=float(dk[4:].mean()),
                                      p=wp(dk))}
    for g, lab_ in [("fmp_human", "SMA FMP (human, n=4)"), ("fmp_mouse", "SMA FMP (mouse, n=6)"), ("dhb", "SMA DHB (mouse, n=2)")]:
        up, kg = med("sma", f"loso_{g}_upper.json"), med("sma", f"loso_{g}_kegg.json")
        b_, k_ = np.mean(list(up.values())), np.mean(list(kg.values()))
        rows.append([lab_, b_, k_, k_ - b_, f"{int(sum(kg[s] > up[s] for s in up))}/{len(up)} sections improve"])
        N["kegg_prior"][g] = dict(base=float(b_), kegg=float(k_), delta=float(k_ - b_))
    cc = A("celltype_ablation_by_class.csv").set_index("class")
    N["kegg_prior"]["by_class"] = {c_: float(cc.loc[c_, "C_+KEGG"] - cc.loc[c_, "B_diffusion"]) for c_ in cc.index}
    B.sheet("kegg_prior", ["cohort", "MetaboCast", "+ KEGG prior", "Δ", "note"], rows,
            ["Held-out-section design, detection-matched median Spearman. The prior: mean z-scored member-gene expression of 85 KEGG "
             "metabolic gene sets (~1,700 human genes) with one neighbour-diffused copy; all else fixed. SMA: gene symbols upper-cased in "
             "both arms so that the human gene sets match mouse symbols.",
             "The released model is trained without the prior (use_kegg=False); use_kegg=True enables it. Class-level changes: " + TR("celltype") + "."],
            [36, 14, 14, 10, 36])

    # ------------------------------------------------------------- cell-type features
    cl = A("celltype_ablation_persection.csv").set_index(A("celltype_ablation_persection.csv").columns[0])
    rows = [[s] + list(x.values) for s, x in cl.iterrows()] + [["mean over sections"] + list(cl.mean().values)]
    rows += [[], ["Per putative class (mean per-ion Spearman over held-out sections)", "n ions"] + list(cl.columns)]
    rows += [[c_, int(x.n_ions)] + [x[k] for k in cl.columns] for c_, x in cc.iterrows()]
    N["celltype"] = {k: float(x) for k, x in cl.mean().items()}
    N["celltype_wins"] = {k: int((cl[k] > cl[cl.columns[0]]).sum()) for k in cl.columns[1:]}
    N["celltype_tests"] = {k: dict(delta=boot_ci(cl[k].values, cl[cl.columns[0]].values)[0],
                                   ci=boot_ci(cl[k].values, cl[cl.columns[0]].values)[1:],
                                   wins=int((cl[k] > cl[cl.columns[0]]).sum()), n=len(cl), p=wp(cl[k].values - cl[cl.columns[0]].values))
                           for k in cl.columns[1:]}
    N["celltype_class"] = {c_: {k: float(x[k]) for k in cl.columns} for c_, x in cc.iterrows()}
    head18 = {"B_diffusion": "B: diffusion features", "C_+KEGG": "C: + KEGG prior", "D_+celltype": "D: + cell-type signatures",
              "E_+celltype_nbhd": "E: D + neighbour-averaged copy"}
    rows[len(cl) + 2] = rows[len(cl) + 2][:2] + [head18.get(k, k) for k in cl.columns]
    A_ = pd.read_csv(os.path.join(R, "desium_ion_annotation.csv")); nm_ = A_.annotation.fillna("").astype(str)
    rx = A_[((A_.n_candidates == 1) & nm_.str.lower().str.startswith(("glutathione", "gssg", "ascorbate", "urate", "taurine")))
            | (A_["class"] == "peptide/redox")]
    rx_txt = ", ".join(f"{n_} {a_.replace('-', '−')}" for n_, a_ in zip(rx.annotation, rx.adduct))
    B.sheet("celltype", ["section"] + [head18.get(k, k) for k in cl.columns], rows,
            ["B: diffusion features; C: + KEGG prior; D: + reference-free per-spot lineage-signature block via extra_key; E: D + its "
             "neighbour-averaged copy. Held-out-section design, n_hvg = 2,000 (reproduce/celltype_ablation.py).",
             "Classes: channels whose putative candidates share one class (reproduce/annotate_ions.py); 'redox' = channels with a single "
             "candidate named glutathione, GSSG, ascorbate, urate or taurine, or of class peptide/redox (here " + rx_txt + "), which "
             "can also count under their chemical class. Each mean covers the held-out sections where the ion was detected in >20% of "
             "spots."], [44, 14, 14, 16, 18])

    # ------------------------------------------------------------- robustness source values
    RG = lambda n_: J("rigor", f"{n_}.json")
    rb, rc, rs, rg, ru, rsc, rsh = (RG("rigor_baselines"), RG("rigor_ceiling"), RG("rigor_splits"), RG("rigor_gcn"), RG("rigor_uncertainty"),
                                    RG("rigor_spatial_control"), RG("rigor_shift"))
    rp = RG("rigor_predictability")
    rows = [["Extended baselines (1,000 ions, 3 sections, spatial 3-fold)", k, x[0], x[1]] for k, x in rb["baselines"].items()]
    rows += [[f"Learner family ({FR('SuppFig_robustness')}a; {rc['n_metab']} ions)", k, x[0], x[1]] for k, x in rc["family"].items()]
    rows += [[f"PCA dimensionality ({rc['n_metab']} ions)", f"{k} PCs", x, None] for k, x in rc["pca_dim"].items()]
    rows += [[f"Diffusion depth ({FR('SuppFig_robustness')}b; {rs['n_metab']} ions)", f"{k} hops", x[0], x[1]] for k, x in rs["hops"].items()]
    rows += [[f"Validation split ({FR('SuppFig_robustness')}c; {rs['n_metab']} ions)", k, x[0], x[1]] for k, x in rs["splits"].items()]
    rows.append([f"Validation split ({FR('SuppFig_robustness')}c)", "leave-region-out", rb["leave_region_out"][0], rb["leave_region_out"][1]])
    rows += [[f"Abundance quintile ({FR('SuppFig_robustness')}d)", f"Q{i} (mean log abundance {ab:.2f})", x, None] for i, (ab, x) in enumerate(rb["abundance"], 1)]
    det_ = f"Determinants of predictability ({FR('SuppFig_predictability')}a,b; ranking in {TR('robustness')}b)"
    rows += [[det_, f"univariate Spearman with {k}", x, None] for k, x in rp["univariate"].items()]
    rows += [[det_, f"permutation importance (Δ held-out R²): {k}", x["loso_mean"], None] for k, x in rp["permutation"].items()]
    rows += [[det_, "held-out R², leave-one-section-out (4 folds)", rp["cv_r2"], None],
             [det_, "held-out R², five folds of ions", rp["cv_r2_ion_folds"], None]]
    rows += [[f"Fixed diffusion vs learned GCN, both on rank targets ({FR('SuppFig_predictability')}e)", k, x[0], x[1]] for k, x in rg.items()]
    shown_ = {"per-spot ridge (no graph)": "no diffusion"}         # result key -> label used in the text and figures
    rows += [[f"Spatial-structure controls ({FR('SuppFig_spatial_controls')})", shown_.get(k, k), x[0], x[1]] for k, x in rsc["conditions"].items()]
    rows += [[f"Uncertainty under shift ({FR('SuppFig_robustness')}f; 200 ions)", "coverage, in-distribution", ru["coverage_id"], None],
             ["Uncertainty under shift", f"coverage, new section ({ru['held']})", ru["coverage_ood"], None],
             ["Uncertainty under shift", "epistemic ratio new / in-distribution", ru["epistemic_ratio"], None]]
    rows += [[f"Coverage by epistemic-uncertainty quartile ({FR('SuppFig_predictability')}d)", f"Q{q} (centre {c_:.3f})", cv_, None]
             for q, (c_, cv_) in enumerate(rsh["epistemic_calibration"], 1)]
    rows += [[f"Transfer vs distribution shift ({FR('SuppFig_predictability')}c)", "Spearman(metabolite shift, held-out r)", rsh["spearman_metab_loso"], None],
             ["Transfer vs distribution shift", "Spearman(transcriptome shift, held-out r)", rsh["spearman_embed_loso"], None]]
    pg = J("penalty_grid.json"); lab_ = "Ridge penalty beyond the default grid (10⁻¹–10⁷ extended to 10¹⁰; reproduce/penalty_grid_check.py)"
    gi_ = lambda c, a: c["mean"][c["alpha"].index(a)]
    rows += [[lab_, "shipped configuration: penalty chosen, default / extended grid",
              f"{pw(pg['shipped']['chosen_default'])} / {pw(pg['shipped']['chosen_extended'])}", None],
             [lab_, "shipped configuration: inner CV score at the two penalties",
              f"{gi_(pg['shipped'], pg['shipped']['chosen_default']):.4f} / {gi_(pg['shipped'], pg['shipped']['chosen_extended']):.4f}", None],
             [lab_, "held-out-section accuracy with the penalty of the default / extended grid (mean of 7 sections)",
              f"{pg['transfer']['mean_default']:.4f} / {pg['transfer']['mean_extended']:.4f}", None],
             [lab_, "held-out sections whose selected penalty changed",
              f"{sum(v['chosen_default'] != v['chosen_extended'] for v in pg['transfer']['per_section'].values())} of 7", None],
             [lab_, "HCC/GBM application model: penalty chosen, default / extended grid",
              f"{pw(pg['application']['chosen_default'])} / {pw(pg['application']['chosen_extended'])}", None],
             [lab_, "HCC/GBM application model: inner CV score at the two penalties",
              f"{gi_(pg['application'], pg['application']['chosen_default']):.4f} / "
              f"{gi_(pg['application'], pg['application']['chosen_extended']):.4f}", None]]
    N["penalty_grid"] = dict(shipped=[pg["shipped"]["chosen_default"], pg["shipped"]["chosen_extended"]],
                             transfer=[pg["transfer"]["mean_default"], pg["transfer"]["mean_extended"]],
                             n_changed=sum(v["chosen_default"] != v["chosen_extended"] for v in pg["transfer"]["per_section"].values()),
                             application=[pg["application"]["chosen_default"], pg["application"]["chosen_extended"]])
    gm = J("gbm_multimodal.json")
    rows += [["GBM CytAssist proteins", "mean RNA-reconstructable variance (R²)", gm["mean_r2"], None],
             ["GBM CytAssist proteins", "mean variance orthogonal to RNA", gm["mean_orthogonal"], None]]
    N["robustness"] = dict(predictability=dict(univariate=rp["univariate"], univariate_ci=rp["univariate_ci"], ranking=rp["ranking"],
                                               permutation=rp["permutation"], impurity=rp["impurity"], cv_r2=rp["cv_r2"],
                                               cv_r2_ion_folds=rp["cv_r2_ion_folds"], fold_r2_sections=rp["fold_r2_sections"],
                                               n=rp["n"], n_ions=rp["n_ions"], n_by_group=rp["n_by_group"]),
                           gbm_protein=dict(mean_r2=gm["mean_r2"], orthogonal=gm["mean_orthogonal"]))
    B.sheet("robustness", ["analysis", "condition", "median Spearman / value", "s.d. across sections/folds"], rows,
            ["Smaller well-measured ion panels (n ions given per row), so absolute values are not comparable with the top-1,000 benchmark "
             f"of {TR('benchmark')}. MetaboCast rows use the released estimator (reproduce/estimator.py).",
             "Learner-family rows other than the two MetaboCast rows are fitted to log intensities; the GCN is trained on the same "
             "rank-Gaussian targets as MetaboCast.",
             "Determinants of predictability (reproduce/rigor_predictability.py): per-ion within-section r on four sections regressed on "
             "abundance, coefficient of variation, Moran's I, detection rate and putative chemical group (lipid, polar or unannotated); "
             f"gradient boosting, scored on held-out sections and on held-out ions; ranking and intervals in {TR('robustness')}b.",
             "Transfer vs distribution shift: Spearman correlations over seven sections; neither is statistically significant.",
             "Ridge penalty: the penalty is chosen by the median per-ion Pearson correlation of inner held-out sections, which does not "
             "depend on the scale of the predictions and changes little above 10⁷; larger penalties mainly shrink the ridge half and, "
             f"through the equal-spread rule, the amplitude of the ensemble. HCC contrasts at both application penalties: {TR('hcc_patients')}b."],
            [56, 46, 22, 22])

    # ------------------------------------------------------------- model selection
    rows = []
    keysL = [("pca|a200", "ridge, fixed alpha 200, log targets"), ("pca|aG", "ridge, CV penalty, log targets"),
             ("pca|yR|aG", "ridge, CV penalty, rank targets"), ("pca|yR|bag", "jackknife-bagged ridge, rank targets"),
             ("pca|yR|ET", "ExtraTrees, rank targets"), ("pca|yR|ens(ridge aG,ET)", "ridge + ExtraTrees, rank targets (released)"),
             ("secz|yR|aG", "section z-scored genes, ridge"), ("secz|yR|ens(ridge aG,ET)", "section z-scored genes, ensemble"),
             ("secrank|yR|aG", "section rank genes, ridge"), ("pca+enz|yR|aG", "PCA + enzyme-gene block, ridge"), ("pca200|yR|aG", "PCA-200, ridge")]

    def add(label, rel_path, keys):
        ps = J("model_selection", rel_path)["per_section"]
        for k, lab2 in keys:
            if k in ps:
                vv = list(ps[k].values()); rows.append([label, lab2, k, float(np.mean(vv)), len(vv)])
    add("DESIUM held-out section (7)", "r3_loso_desium.json", keysL)
    add("DESIUM held-out section (7)", "loso2_desium.json", [("nmf|yR|aG", "NMF-100 gene representation, ridge, rank targets")])
    for g, lab2 in [("dhb", "SMA DHB held-out section (2)"), ("fmp_human", "SMA FMP human held-out section (4)"),
                    ("fmp_mouse", "SMA FMP mouse held-out section (6)")]:
        add(lab2, f"r3_loso_sma_{g}.json", keysL)
    add("DESIUM held-out section (7), TIC-normalised target", "loso_tic_desium.json", [("pca|aG|TICtarget", "TIC-normalised target")])
    j = J("model_selection", "r3_within_desium.json")
    for k, lab2 in [("pca|y|a200", "ridge, fixed alpha 200 (intensity scale)"), ("pca|y|aG", "ridge, CV penalty, log targets (intensity scale)"),
                    ("pca|yR|aG", "ridge, rank targets (rank scale)"), ("pca|yR|ET", "ExtraTrees, rank targets (rank scale)"),
                    ("ens(pca|yR|aG,ET)", "ridge + ExtraTrees (rank scale)"), ("pca+enz|yR|aG", "PCA + enzyme block (rank scale)")]:
        vv = [j[s]["scores"][k] for s in j if k in j[s]["scores"]]
        if vv:
            rows.append(["DESIUM within-section (7), screen", lab2, k, float(np.mean(vv)), len(vv)])
    combo_lab = {"allions|ridge": "ridge", "trees": "ExtraTrees", "allions|mean": "plain mean", "allions|eqSD": "equal-spread mean (released)",
                 "panel1000|eqSD": "equal-spread, penalty chosen on the 1,000-ion panel"}
    for name, lab2 in [("combo_loso_desium.json", "DESIUM held-out section (7)"), ("combo_loso_dhb.json", "SMA DHB held-out section (2)"),
                       ("combo_loso_fmph.json", "SMA FMP human held-out section (4)")]:
        j = J("model_selection", name)
        for k, l3 in combo_lab.items():
            vv = [j[s][k] for s in j if k in j[s]]
            rows.append([lab2 + ", combination rule (package protocol)", l3, k, float(np.mean(vv)), len(vv)])
    for f, lab2 in [("screen_alpha_sma.json", "SMA within-section, single-section penalty rule"),
                    ("screen_alpha_desium.json", "DESIUM within-section, single-section penalty rule")]:
        j = J("model_selection", f)
        for k, l3 in [("perblock|ens|log", "inner blocks scored one at a time"), ("pooled|ens|log", "inner held-out blocks pooled (released rule)"),
                      ("a200|ens|log", "fixed alpha = 200"), ("a3e4|ens|log", "fixed alpha = 30,000"),
                      ("fixed200_logtargets|log", "ridge, alpha 200, log targets, no trees")]:
            vv = [j[s]["scores"][k] for s in j]
            rows.append([f"{lab2} ({len(vv)} sections; intensity scale)", l3, k, float(np.mean(vv)), len(vv)])
        N[f"screen_{f.split('_')[-1].split('.')[0]}"] = {k: float(np.mean([j[s]["scores"][k] for s in j])) for k in ("perblock|ens|log", "pooled|ens|log")}
        N[f"screen_{f.split('_')[-1].split('.')[0]}"]["n"] = len(j)
    N["model_selection"] = {lab: {r_[2]: r_[3] for r_ in rows if r_[0] == lab} for lab in dict.fromkeys(r_[0] for r_ in rows)}
    B.sheet("model_selection", ["setting", "variant", "code", "mean median Spearman", "n sections"], rows,
            ["Candidates came from published approaches: rank-based harmonisation of metabolite targets (MIRTH), strong shrinkage over "
             "genome-wide features, tree ensembles, enzyme-gene and cell-composition features, and per-section scaling. Every setting "
             "was chosen by inner cross-validation on training data only; choices were made on DESIUM and then checked on the three SMA "
             "cohorts with the published SMA registration (reproduce/model_selection).",
             "Screen held-out-section rows use the screen implementation (pooled 2,000 HVGs for DESIUM); released-package values are in "
             f"{TR('transfer')} and {TR('sma_transfer')}. Within-section screen rows are scored on each variant's training-target scale and are not comparable "
             "with intensity-scale rows.",
             "Per-section gene z-scoring is an option (gene_scaling='section'), not the default: its gain was inconsistent across sections "
             "and it removes between-section differences.",
             "Single-section penalty rule: with one training section the penalty is chosen by leave-one-spatial-block-out CV, pooling the "
             "held-out blocks into one map before scoring (the released rule)."], [48, 46, 30, 16, 10])

    # ------------------------------------------------------------- predictability: feature ranking
    labp = {"moran": "Moran's I (spatial autocorrelation of the measured ion)", "cv": "coefficient of variation",
            "abundance": "mean log abundance", "detection": "detection rate",
            "chemical group": "putative chemical group (lipid, polar or unannotated)"}
    rows = []
    for k_, g_ in enumerate(rp["ranking"], 1):
        pm_ = rp["permutation"][g_]
        imp_ = rp["impurity"][g_] if g_ in rp["impurity"] else rp["impurity"]["is_lipid"] + rp["impurity"]["is_polar"]
        ci_ = rp["univariate_ci"].get(g_)
        rows.append([k_, labp[g_], pm_["loso_mean"], f"{pm_['loso_min']:.3f} to {pm_['loso_max']:.3f}", pm_["ion_folds_mean"], imp_,
                     rp["univariate"].get(g_, "n/a"), f"{ci_[0]:+.3f} to {ci_[1]:+.3f}" if ci_ else "n/a"])
    fr_ = rp["fold_r2_sections"]
    B.sheet_b("robustness", "Determinants of predictability: feature ranking",
            ["Rank", "Property", "Decrease in held-out R² when permuted (mean over 4 held-out sections)", "Range over held-out sections",
             "Decrease in held-out R² (5 folds of ions)", "Impurity importance (model fitted to all rows)",
             "Univariate Spearman ρ with accuracy", "95% CI of ρ (bootstrap over ions)"], rows,
            [f"Unit: ion-section pairs (n = {rp['n']:,}; {rp['n_ions']:,} ions detected in >20% of spots of the DESIUM sections "
             f"{', '.join(rp['sections'])}). Outcome: per-ion Spearman r of five-fold spatial-block cross-validated predictions "
             "(released estimator). Groups: " + ", ".join(f"{k} {v:,}" for k, v in rp["n_by_group"].items()) + " pairs.",
             f"Model: gradient boosting (200 trees, depth 3, learning rate 0.05). Held-out R² {rp['cv_r2']:.3f} for the pooled out-of-fold "
             f"predictions leaving out one section at a time (per held-out section: {', '.join(f'{s} {v:.2f}' for s, v in fr_.items())}) "
             f"and {rp['cv_r2_ion_folds']:.3f} over five folds of ions "
             "(no ion in both training and test rows).",
             "Ranking: decrease in held-out R² when one property is permuted within the held-out fold (20 permutations per fold; the "
             "lipid and polar indicators are permuted together); values above the held-out R² mean that the permuted model is worse "
             "than predicting the mean. Impurity importance is shown for reference only (fitted to all rows, favours continuous "
             "features). Univariate intervals resample ions, keeping all sections of an ion together (2,000 resamples).",
             "Source: reproduce/rigor_predictability.py."], [6, 42, 20, 16, 18, 18, 16, 18])

    # ------------------------------------------------------------- statistics register
    sr = []                                            # family, unit, n, comparison, delta, ci, improved, test, p
    WX = "exact two-sided Wilcoxon signed-rank"
    nice = {"joint-NMF": "joint NMF", "coordinates only (tuned kNN)": "coordinate-only kNN", "no diffusion": "no diffusion",
            "KEGG pathway features + supervised regressor": "KEGG pathway features + regressor",
            "MetaboCast (ridge only)": "MetaboCast, ridge only", "single-gene OLS": "single-gene regression"}
    for key_, d_ in N["benchmark"].items():
        a_, b_ = key_.split(" - ")
        sr.append(["A. Within-section benchmark, DESIUM", "section", d_["n"], f"{nice.get(a_, a_)} − {nice.get(b_, b_)}",
                   d_["delta"], d_["ci"], f"{d_['wins']}/{d_['n']}", WX, d_["p"]])
    for k_, lab_ in [("vs ridge only", "MetaboCast, ridge only"), ("vs graph-free", "no diffusion"),
                     ("vs KEGG features", "KEGG pathway features + regressor"), ("vs joint NMF", "joint NMF"),
                     ("vs single gene", "single-gene regression")]:
        if k_ in N["transfer"]:
            d_ = N["transfer"][k_]
            sr.append(["B. Held-out-section transfer, DESIUM", "section", 7, f"MetaboCast − {lab_}", d_["delta"], d_["ci"],
                       f"{d_['wins']}/7", WX, d_["p"]])
    kp_ = N["kegg_prior"]["DESIUM"]
    sr.append(["B. Held-out-section transfer, DESIUM", "section", 7, "MetaboCast + KEGG prior − MetaboCast", kp_["delta"], kp_["ci"],
               f"{kp_['wins']}/7", WX, kp_["p"]])
    for k_, lab_ in [("D_+celltype", "MetaboCast + cell-type signatures − MetaboCast")]:
        if k_ in N["celltype_tests"]:
            d_ = N["celltype_tests"][k_]
            sr.append(["B. Held-out-section transfer, DESIUM", "section", d_["n"], lab_, d_["delta"], d_["ci"], f"{d_['wins']}/{d_['n']}",
                       WX, d_["p"]])
    h_ = N["histology"]
    sr.append(["C. Histology, held-out-section transfer, DESIUM", "section", h_["n"], "H&E only − transcriptome only",
               h_["he_only"] - h_["transcriptome"], h_["ci_he"], f"{h_['he_wins']}/{h_['n']}", WX, h_["p_he"]])
    sr.append(["C. Histology, held-out-section transfer, DESIUM", "section", h_["n"], "transcriptome + H&E − transcriptome only",
               h_["combined"] - h_["transcriptome"], h_["ci_combined"], f"{h_['combined_wins']}/{h_['n']}", WX, h_["p_combined"]])
    st_ = N["sma_within_tests"]
    for k_, lab_ in [("fmp_vs_nmf", "MetaboCast − joint NMF, FMP sections"), ("fmp_vs_graphfree", "MetaboCast − no diffusion, FMP sections"),
                     ("all_vs_nmf", "MetaboCast − joint NMF, all SMA sections")]:
        d_ = st_[k_]
        sr.append(["D. Within-section benchmark, SMA", "section", d_["n"], lab_, d_["delta"], d_["ci"], f"{d_['wins']}/{d_['n']}", WX, d_["p"]])
    for tool in ("scFEA", "MGFEA"):
        d_ = N["flux_baselines"][tool]
        sr.append(["E. Held-out MetaboCast vs transcriptome-only inference", "ion", d_["n"],
                   f"MetaboCast (held-out section) − {tool}, median paired difference", d_["loso_median_diff"][0],
                   d_["loso_median_diff"][1:], f"{d_['loso_wins']}/{d_['n']}", WX,
                   d_["p_loso"]])
    sl_ = rsh["shift_loso"]
    for j_, lab_ in [(3, "metabolite distribution shift"), (2, "transcriptome distribution shift")]:
        rho_, p_ = spearmanr([r_[j_] for r_ in sl_], [r_[1] for r_ in sl_])
        sr.append(["F. Transfer accuracy vs distribution shift, DESIUM", "section", len(sl_), f"Spearman ρ, {lab_}",
                   float(rho_), None, "", "Spearman correlation test", float(p_)])
    fams = list(dict.fromkeys(r_[0] for r_ in sr)); adj_ = {}
    for f_ in fams:
        idx_ = [i for i, r_ in enumerate(sr) if r_[0] == f_]
        for i, a_ in zip(idx_, holm([sr[i][8] for i in idx_])):
            adj_[i] = float(a_)
    rows = [[r_[0], r_[3], r_[1], r_[2], r_[4], (f"{r_[5][0]:+.4f} to {r_[5][1]:+.4f}" if r_[5] is not None else "n/a"), r_[6], r_[7],
             r_[8], adj_[i], sum(1 for x in sr if x[0] == r_[0])] for i, r_ in enumerate(sr)]
    N["statistics"] = [dict(family=r_[0], comparison=r_[1], unit=r_[2], n=r_[3], effect=r_[4], ci=r_[5], improved=r_[6], test=r_[7],
                            p=r_[8], p_holm=r_[9], family_size=r_[10]) for r_ in rows]
    B.sheet("statistics", ["Family", "Comparison", "Unit", "n", "Effect", "95% CI (paired bootstrap, 5,000)", "Units improved", "Test",
                           "P", "P, Holm-adjusted within family", "Tests in family"], rows,
            ["Every inferential comparison reported in the manuscript. Effects are paired differences of median per-ion Spearman "
             "correlations (families A-D), the median paired difference over ions (E) or a Spearman correlation (F); intervals are "
             "percentile intervals from 5,000 bootstrap resamples of the paired units (sections or ions).",
             "P values are adjusted with Holm's step-down procedure within each family (family-wise error rate). With seven sections the "
             "smallest attainable exact two-sided Wilcoxon P is 0.016, so no seven-section comparison can remain below 0.05 after "
             "adjustment over several comparisons; the paired effects, their intervals and the number of sections improved are the "
             "primary evidence, and the tests serve as consistency checks.",
             "Descriptive results without tests: coverage of the uncertainty widths (" + TR("coverage") + "), the SMA transfer estimates "
             "(" + TR("sma_transfer") + "; one to three animals or one donor per cohort), and the HCC and GBM predictions, which have no "
             "measured metabolites."], [44, 52, 8, 5, 10, 22, 10, 30, 10, 14, 10])

    # ------------------------------------------------------------- README sheet and order
    ws = B.wb.create_sheet("README", 0)
    init = open(os.path.join(os.path.dirname(os.path.dirname(HERE)), "metabocast", "__init__.py")).read()
    version = re.search(r'__version__ = "([^"]+)"', init).group(1)
    ws.append([f"Additional file 2: Supplementary Tables S1-S{len(TABLES)} (MetaboCast {version})"]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Sheet", "Contents"])
    for c in ws[2]:
        c.font = Font(bold=True); c.fill = HFILL
    for t in TABLES:
        if table_sheet(t["key"]) in B.wb.sheetnames:
            ws.append([table_sheet(t["key"]), f"{t['label']}: {t['title']}"])
            if t["key"] == "per_ion":
                ws.append([table_sheet("per_ion") + "b", f"Table {table_sheet('per_ion')}b: All 2,086 DESIUM m/z channels: accuracy, tier and "
                                                    "putative candidates"])
            elif table_sheet(t["key"]) + "b" in B.b_titles:
                ws.append([table_sheet(t["key"]) + "b", f"Table {table_sheet(t['key'])}b: {B.b_titles[table_sheet(t['key']) + 'b']}"])
    ws.append([]); ws.append(["Every table is rebuilt from reproduce/results by reproduce/tables/build_supplementary_tables.py "
                              "(https://github.com/himanshushekhar367-sudo/MetaboCast)."])
    ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 110
    order = ["README"]
    for t in TABLES:
        order.append(table_sheet(t["key"]))
        if t["key"] == "per_ion" or table_sheet(t["key"]) + "b" in B.b_titles:
            order.append(table_sheet(t["key"]) + "b")
    B.wb._sheets = [B.wb[n] for n in order if n in B.wb.sheetnames] + [s for s in B.wb._sheets if s.title not in order]
    B.wb.save(out)
    json.dump(N, open(os.path.splitext(out)[0] + "_notes.json", "w"), indent=1,
              default=lambda o: list(o) if isinstance(o, tuple) else (o.tolist() if hasattr(o, "tolist") else float(o)))
    print("saved", out, "|", ", ".join(B.wb.sheetnames))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--results", default=os.path.join(os.path.dirname(HERE), "results"))
    ap.add_argument("--out", default="MetaboCast_Supplementary_Tables.xlsx")
    args = ap.parse_args()
    build(args.results, args.out)


if __name__ == "__main__":
    main()
